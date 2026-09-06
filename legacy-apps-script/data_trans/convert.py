"""Read the legacy workbook without modifying it; produce JSON for import.mjs."""

import argparse
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation
import hashlib
import json
from pathlib import Path
import re
import sys

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
TAIPEI = timezone(timedelta(hours=8))
SCHEMAS = {
    "Products": ["ProductID", "ProductName", "Category", "Price", "Status", "CreateTime",
                 "Description", "GiftBoxEnabled", "SpecialPrice", "CompanyPrice"],
    "Customers": ["CustomerID", "Name", "Phone", "Address", "CreateTime", "LastOrderDate"],
    "Orders": ["OrderID", "CustomerName", "CustomerPhone", "CustomerAddress", "DeliveryType",
               "DeliveryDate", "Status", "CreateTime", "TotalAmount", "DepositAmount",
               "RemainingAmount", "PaymentNotes", "ShippingFee", "ShippingNotes", "Notes",
               "RecipientName", "RecipientPhone", "IsCompanyCustomer"],
    "OrderDetails": ["DetailID", "OrderID", "ProductID", "ProductName", "Quantity", "UnitPrice",
                     "Subtotal", "IsGiftBox", "GiftBoxDetails", "OriginalPrice", "IsSpecialPrice"],
    "CapacitySettings": ["ID", "Type", "DayOfWeek", "Date", "MaxQuantity", "Enabled",
                         "CreateTime", "UpdateTime"],
}
ALIASES = {
    "Orders": {"Address": "CustomerAddress", "DeliverType": "DeliveryType", "NA": "Notes",
               "ReceivedName": "RecipientName", "ReceivedAddress": "RecipientPhone"},
    "CapacitySettings": {"CAPid": "ID", "weekday": "Type", "weekday_no": "DayOfWeek",
                         "quantity": "MaxQuantity", "enable": "Enabled",
                         "generate_date": "CreateTime", "modify_date": "UpdateTime"},
}
REQUIRED = {
    "Products": {"ProductID", "ProductName", "Price", "Status", "CreateTime"},
    "Customers": {"CustomerID", "Name", "Phone", "CreateTime"},
    "Orders": {"OrderID", "CustomerName", "CustomerPhone", "DeliveryDate", "Status",
               "CreateTime", "TotalAmount"},
    "OrderDetails": {"DetailID", "OrderID", "ProductID", "ProductName", "Quantity",
                     "UnitPrice", "Subtotal", "IsGiftBox", "GiftBoxDetails"},
    "CapacitySettings": {"Type", "DayOfWeek", "MaxQuantity", "Enabled"},
}


def clean(value):
    if isinstance(value, str):
        return value[1:] if value.startswith("'") else value
    return value


def text(value):
    return "" if value is None else str(clean(value)).strip()


def blank(value):
    return text(value) == ""


def number(value, default=None, integer=False):
    if blank(value):
        if default is None:
            raise ValueError("缺少數字")
        return default
    try:
        result = Decimal(text(value))
    except InvalidOperation as error:
        raise ValueError("無法解析數字") from error
    if not result.is_finite() or result < 0 or result > Decimal("9007199254740991"):
        raise ValueError("數字必須為有效且非負的安全數值")
    if integer and result != result.to_integral_value():
        raise ValueError("數量必須為整數")
    return int(result) if result == result.to_integral_value() else float(result)


def yes(value, default=False):
    if blank(value):
        return default
    raw = text(value).lower()
    if raw in ("是", "true", "1"):
        return True
    if raw in ("否", "false", "0"):
        return False
    raise ValueError("無法解析是／否欄位")


def timestamp(value):
    if blank(value):
        raise ValueError("缺少建立時間")
    if isinstance(value, datetime):
        result = value
    elif isinstance(value, date):
        result = datetime.combine(value, datetime.min.time())
    else:
        result = datetime.fromisoformat(text(value).replace("Z", "+00:00"))
    return result.replace(tzinfo=result.tzinfo or TAIPEI).isoformat(timespec="seconds")


def day(value):
    result = datetime.fromisoformat(timestamp(value)).astimezone(TAIPEI)
    return result.date().isoformat()


def document_id(value):
    result = text(value)
    if not result or "/" in result or result in (".", "..") or len(result.encode()) > 1500:
        raise ValueError("無效或空白的資料編號")
    return result


def contact(value, name, fallback, warnings, location):
    # Numeric Excel phone cells may already have lost their leading zero.
    if value is not None and not isinstance(value, str):
        raise ValueError("電話欄為數字型別，請在來源確認前導 0 並改存文字")
    raw = text(value)
    is_line = raw.lower() == "line"
    normalized = "line" if is_line else re.sub(r"\D", "", raw, flags=re.ASCII)
    if is_line:
        key = hashlib.sha256(f"LINE:{name}".lower().encode()).hexdigest()[:32]
        identifier = "line_" + key
    else:
        identifier = normalized or "legacy_" + document_id(fallback)
        if not normalized or not re.fullmatch(r"[0-9+()\s-]+", raw):
            warnings.append({"location": location, "code": "contact_needs_review",
                             "message": "聯絡欄為空白或含非電話文字；保留原文，請核對"})
    return identifier, {"contactType": "line" if is_line else "phone", "contactValue": raw,
                        "contactNormalized": normalized, "phone": "" if is_line else raw,
                        "lineId": raw if is_line else ""}


def read_tables(source):
    workbook = load_workbook(source, read_only=True, data_only=False)
    try:
        tables = {}
        for name in SCHEMAS:
            if name not in workbook.sheetnames:
                raise ValueError(f"缺少工作表 {name}")
            tables[name] = []
            for row in workbook[name].iter_rows():
                values = []
                for cell in row:
                    if cell.data_type in ("f", "e"):
                        raise ValueError(f"{name}!{cell.coordinate}: 公式或錯誤值需先在來源轉成已核對的值")
                    values.append(clean(cell.value))
                tables[name].append(values)
        return tables
    finally:
        workbook.close()


def table_records(name, rows, warnings):
    if not rows:
        raise ValueError(f"{name}: 缺少欄名")
    headers = []
    for i, value in enumerate(rows[0]):
        raw = text(value).lstrip("`")
        header = ALIASES.get(name, {}).get(raw, raw)
        # Known legacy omissions only; named/reordered columns always use their header.
        if not header and (name, i) in {("Orders", 17), ("Products", 9), ("CapacitySettings", 3)}:
            header = SCHEMAS[name][i]
            warnings.append({"location": f"{name}!column {i + 1}", "code": "legacy_header",
                             "message": f"空白欄名依舊版結構映射至 {header}"})
        if header in headers and header:
            raise ValueError(f"{name}: 重複欄名 {header}")
        headers.append(header)
    missing = REQUIRED[name] - set(headers)
    if missing:
        raise ValueError(f"{name}: 缺少欄位 {sorted(missing)}")
    for row_number, row in enumerate(rows[1:], 2):
        if all(blank(v) for v in row):
            continue
        record = {key: row[i] if i < len(row) else None for i, key in enumerate(headers) if key}
        extras = {(headers[i] if i < len(headers) else "") or f"column_{i+1}": value for i, value in enumerate(row)
                  if not blank(value) and (i >= len(headers) or headers[i] not in SCHEMAS[name])}
        if extras:
            raise ValueError(f"{name}!row {row_number}: 有未映射且非空白的欄位 {list(extras)}")
        yield row_number, record


def convert(tables, source_name="database.xlsx", source_hash=""):
    warnings, errors = [], []
    payload = {"format": "gin-jia-pos-xlsx", "schemaVersion": 2,
               "source": {"file": source_name, "sha256": source_hash},
               "products": [], "customers": [], "orders": [],
               "weekdayCapacity": [], "capacityOverrides": []}
    records = {name: list(table_records(name, rows, warnings)) for name, rows in tables.items()}
    counts = {name: len(rows) for name, rows in records.items()}
    details = defaultdict(list)
    seen = defaultdict(set)

    def unique(kind, value):
        key = document_id(value)
        if key in seen[kind]:
            raise ValueError(f"{kind}: 資料編號重複")
        seen[kind].add(key)
        return key

    def warn(location, code, message):
        warnings.append({"location": location, "code": code, "message": message})

    for name in ["Products", "Customers", "OrderDetails", "Orders", "CapacitySettings"]:
        for row_number, r in records[name]:
            loc = f"{name}!row {row_number}"
            try:
                if name == "Products":
                    p = {"productId": unique(name, r.get("ProductID")),
                         "productName": text(r.get("ProductName")), "category": text(r.get("Category")),
                         "price": number(r.get("Price")), "status": text(r.get("Status")),
                         "createTime": timestamp(r.get("CreateTime")), "description": text(r.get("Description")),
                         "giftBoxEnabled": "是" if yes(r.get("GiftBoxEnabled"), True) else "否",
                         "specialPrice": number(r.get("SpecialPrice"), ""),
                         "companyPrice": number(r.get("CompanyPrice"), "")}
                    if not p["productName"] or p["status"] not in ("啟用", "停用"):
                        raise ValueError("商品名稱或狀態無效")
                    payload["products"].append(p)
                elif name == "Customers":
                    legacy_id = unique(name, r.get("CustomerID"))
                    cname = text(r.get("Name"))
                    cid, fields = contact(r.get("Phone"), cname, legacy_id, warnings, loc)
                    payload["customers"].append({"documentId": cid, "customerId": legacy_id,
                        "name": cname, **fields, "address": text(r.get("Address")),
                        "createTime": timestamp(r.get("CreateTime")),
                        "lastOrderDate": timestamp(r.get("LastOrderDate") or r.get("CreateTime"))})
                elif name == "OrderDetails":
                    gift = yes(r.get("IsGiftBox"))
                    composition = None
                    if gift:
                        composition = json.loads(text(r.get("GiftBoxDetails")))
                        if not isinstance(composition, dict) or not isinstance(composition.get("products"), dict):
                            raise ValueError("禮盒缺少有效的內容明細")
                        composition["size"] = number(composition.get("size"), integer=True)
                        composition["products"] = {document_id(k): number(v, integer=True)
                                                   for k, v in composition["products"].items()}
                        if composition["size"] <= 0:
                            raise ValueError("禮盒規格必須大於 0")
                    item = {"detailId": unique(name, r.get("DetailID")),
                            "productId": document_id(r.get("ProductID")), "productName": text(r.get("ProductName")),
                            "quantity": number(r.get("Quantity"), integer=True), "unitPrice": number(r.get("UnitPrice")),
                            "subtotal": number(r.get("Subtotal")), "isGiftBox": gift, "giftBoxDetails": composition,
                            "originalPrice": number(r.get("OriginalPrice"), number(r.get("UnitPrice"))),
                            "isSpecialPrice": yes(r.get("IsSpecialPrice"))}
                    if item["quantity"] <= 0:
                        raise ValueError("商品數量必須大於 0")
                    if abs(item["quantity"] * item["unitPrice"] - item["subtotal"]) > 0.01:
                        warn(loc, "subtotal_mismatch", "數量乘單價與小計不同；保留來源小計")
                    details[document_id(r.get("OrderID"))].append(item)
                elif name == "Orders":
                    oid = unique(name, r.get("OrderID"))
                    cname = text(r.get("CustomerName"))
                    _, fields = contact(r.get("CustomerPhone"), cname, oid, warnings, loc)
                    total, deposit = number(r.get("TotalAmount")), number(r.get("DepositAmount"), 0)
                    remaining = number(r.get("RemainingAmount"), max(0, total - deposit))
                    o = {"orderId": oid, "customerName": cname,
                         "customerContactType": fields["contactType"], "customerContactValue": fields["contactValue"],
                         "customerContactNormalized": fields["contactNormalized"],
                         "customerPhone": fields["phone"], "customerLineId": fields["lineId"],
                         "customerAddress": text(r.get("CustomerAddress")),
                         "deliveryType": text(r.get("DeliveryType")) or "外送",
                         "deliveryDate": day(r.get("DeliveryDate")), "status": text(r.get("Status")),
                         "createTime": timestamp(r.get("CreateTime")), "totalAmount": total,
                         "depositAmount": deposit, "remainingAmount": remaining,
                         "paymentNotes": text(r.get("PaymentNotes")), "shippingFee": number(r.get("ShippingFee"), 0),
                         "shippingNotes": text(r.get("ShippingNotes")), "notes": text(r.get("Notes")),
                         "recipientName": text(r.get("RecipientName")), "recipientPhone": text(r.get("RecipientPhone")),
                         "isCompanyCustomer": yes(r.get("IsCompanyCustomer")), "items": details.get(oid, [])}
                    if o["status"] not in ("已確認", "已付訂金", "已付清", "已付款", "完成", "取消"):
                        raise ValueError("未知訂單狀態")
                    if not o["items"]:
                        raise ValueError("訂單沒有明細")
                    if blank(r.get("RemainingAmount")):
                        warn(loc, "remaining_derived", "空白尾款以總額減訂金補入；未依狀態推測付款金額")
                    if abs(deposit + remaining - total) > 0.01:
                        warn(loc, "payment_mismatch", "訂金加尾款不等於總額；保留來源金額")
                    if o["status"] in ("已付款", "已付清") and remaining:
                        warn(loc, "paid_with_balance", "已付款狀態仍有尾款；保留來源狀態，需人工核對")
                    if abs(sum(i["subtotal"] for i in o["items"]) + o["shippingFee"] - total) > 0.01:
                        warn(loc, "total_mismatch", "明細加運費與總額不同；保留來源總額")
                    payload["orders"].append(o)
                else:
                    kind = text(r.get("Type"))
                    setting = {"maxQuantity": number(r.get("MaxQuantity"), "", integer=True),
                               "enabled": yes(r.get("Enabled"))}
                    if kind == "weekday":
                        weekday = number(r.get("DayOfWeek"), integer=True)
                        if weekday > 6:
                            raise ValueError("星期必須為 0 到 6")
                        unique("weekdayCapacity", str(weekday))
                        payload["weekdayCapacity"].append({"dayOfWeek": weekday, **setting})
                    elif kind == "dateOverride":
                        value = day(r.get("Date"))
                        unique("capacityOverrides", value)
                        payload["capacityOverrides"].append({"date": value, **setting})
                    else:
                        raise ValueError("未知供應量設定類型")
            except (ValueError, TypeError, KeyError) as error:
                errors.append({"location": loc, "message": str(error)})

    orphans = set(details) - seen["Orders"]
    if orphans:
        errors.append({"location": "OrderDetails", "message": f"{len(orphans)} 組明細找不到訂單"})
    customers = {}
    for customer in sorted(payload["customers"], key=lambda c: (c["lastOrderDate"], c["customerId"])):
        key = customer["documentId"]
        if key in customers:
            warn("Customers", "normalized_contact_duplicate", "同一標準化聯絡方式合併為最近資料")
        customers[key] = customer
    payload["customers"] = list(customers.values())
    report = {"source": payload["source"], "sourceCounts": counts,
              "outputCounts": {key: len(payload[key]) for key in ("products", "customers", "orders",
                                "weekdayCapacity", "capacityOverrides")},
              "totals": {key: float(sum(Decimal(str(o[key])) for o in payload["orders"]))
                         for key in ("totalAmount", "depositAmount", "remainingAmount", "shippingFee")},
              "warningCounts": dict(Counter(w["code"] for w in warnings)), "warnings": warnings, "errors": errors}
    return payload, report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", nargs="?", type=Path, default=ROOT.parent / "database.xlsx")
    parser.add_argument("--out", type=Path, default=ROOT / "output")
    args = parser.parse_args()
    args.out.mkdir(parents=True, exist_ok=True)
    # Each conversion has its own file; a failed conversion cannot leave a stale import file under that name.
    digest = hashlib.sha256(args.source.read_bytes()).hexdigest()
    output = args.out / f"{args.source.stem}-{digest[:12]}.json"
    report_file = args.out / f"{args.source.stem}-{digest[:12]}.report.json"
    try:
        payload, report = convert(read_tables(args.source), args.source.name, digest)
    except ValueError as error:
        payload, report = None, {"errors": [{"message": str(error)}]}
    report_file.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    if report["errors"]:
        output.write_text(json.dumps({"format": "conversion-failed", "report": str(report_file)}), encoding="utf-8")
        print(f"Conversion failed: {len(report['errors'])} errors. See {report_file}", file=sys.stderr)
        return 1
    output.write_text(json.dumps(payload, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")
    print(json.dumps({"output": str(output), "report": str(report_file), "counts": report["outputCounts"],
                      "warningCounts": report["warningCounts"]}, ensure_ascii=True, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
